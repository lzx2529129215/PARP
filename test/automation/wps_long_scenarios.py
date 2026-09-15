"""AgentNet-inspired long tasks: offline preparation, GUI replay, read-only audits.

These are adapted task chains, not literal human-session recordings. No original
AgentNet code is executed and no GUI result is manufactured by file writers.
"""
import hashlib
import html
import io
import json
from pathlib import Path
import re
import shlex
import subprocess
import time
import xml.etree.ElementTree as ET
import zipfile

from .wps_replay import EngineBackend, guard, snapshot

SOURCE = Path('/home/lzx/Desktop/PARP/test/wps_operation_dataset/outputs/agentnet_sequences_v1/provisional_sequences.json')
TASKS = {
    'word_report': ('20241005204819', '文档制作、保存、PDF 导出与阅读', 'wps', '.docx'),
    'ppt_cover': ('20241014161628', '演示封面文字与对象格式调整', 'wpp', '.pptx'),
    'web_sales': ('20241006132748', '网页隐藏销量检查、录入与汇总', 'et', '.xlsx'),
}
PARAGRAPHS = [('Office Workflow Report', 'Heading1'), ('Overview', 'Heading2'),
    ('This report follows an adapted AgentNet office workflow.', 'Normal'),
    ('Data Collection', 'Heading1'), ('Product sales and units are checked before calculation.', 'Normal'),
    ('Results', 'Heading2'), ('The document is saved and exported for review.', 'Normal')]
PRODUCTS = [dict(id=f'album-{i+1}', name=f'Album version {i+1}', totalSaleAmount=n, units=u)
            for i,(n,u) in enumerate([(43,2),(200,1),(57,3),(299,1),(239,2),(102,1)])]
NS = {'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
PROFILE = dict(calibrated=False, resolution=[1920,1080], canvas_blank=[1100,800],
    pdf_path_label=['Save to','Save path','保存到','保存路径'], pdf_path_offset=[220,0],
    fill_swatch=[420,230], line_swatch=[420,230])


def word_fixture():
    """Use the blank package saved through WPS GUI (control_v4 provenance)."""
    return (Path(__file__).resolve().parents[1] / 'samples/wps/word_long_blank_gui.docx').read_bytes()


def ppt_fixture():
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    p=Presentation();p.slide_width=Inches(13.333);p.slide_height=Inches(7.5)
    slide=p.slides.add_slide(p.slide_layouts[6])
    for name,text,left,top,width,size in [
        ('title','Group assignment report',1,1.1,11,60),('logo','YOUR LOGO',.4,.2,3,20),
        ('description','Delete this description',1,2.7,10,24),
        ('footer','Course | Group',1,6.1,10,22),('group','Reporting Group: Group 2',1,4.2,6,24)]:
        shape=slide.shapes.add_textbox(Inches(left),Inches(top),Inches(width),Inches(.7));shape.name=name
        shape.text=text;shape.text_frame.paragraphs[0].runs[0].font.size=Pt(size)
        if name=='group':
            shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string('FFD8B0')
            shape.line.color.rgb=RGBColor.from_string('000000')
    out=io.BytesIO();p.save(out);return out.getvalue()


def sheet_fixture():
    import openpyxl
    w=openpyxl.Workbook();w.active.title='Sales'
    w.active.append(['Product','Sales','Units','Product amount'])
    out=io.BytesIO();w.save(out);w.close();return out.getvalue()


def task_steps(kind):
    steps=[]
    def add(op,**args):steps.append(dict(id=f'{kind}.{len(steps)+1:03}',op=op,**args))
    if kind=='word_report':
        add('folder_browse');add('open_document')
        for text,style in PARAGRAPHS:add('word_paragraph',text=text,style=style)
        add('save');add('audit_word');add('export_pdf');add('audit_pdf');add('read_pdf');add('close_apps')
    elif kind=='ppt_cover':
        add('open_document');add('shape_text',index=1,text='Why People Obey Authority')
        for size in (56,48,52):add('shape_font',index=1,size=size)
        add('shape_delete',index=2);add('shape_delete',index=2)
        add('shape_text',index=2,text='Mind and Society | Second Group Presentation')
        add('shape_duplicate',index=3)
        add('shape_text',index=4,text='Report Date: 24.10.15')
        add('shape_fill',index=4);add('shape_outline',index=4)
        add('shape_font',index=2,size=20);add('save');add('audit_ppt');add('close_apps')
    elif kind=='web_sales':
        add('open_document')
        for i in range(len(PRODUCTS)):
            add('inspect_product',product_index=i);add('record_product',product_index=i,row=i+2)
        add('sales_formulas');add('save');add('audit_sales');add('close_apps')
    else:raise ValueError(kind)
    return steps


def build(relative='outputs/executable_wps_long_v1'):
    g=guard();g.check();root=g.path(relative)
    if root.exists():raise FileExistsError('Build to a fresh directory')
    sources=json.loads(SOURCE.read_text());entries=[];hashes={}
    def put(name,data):
        with g.open(relative+'/'+name) as f:f.write(data)
        hashes[name]=hashlib.sha256(data).hexdigest()
    def js(name,value):put(name,(json.dumps(value,ensure_ascii=False,indent=2)+'\n').encode())
    for kind,(prefix,title,exe,ext) in TASKS.items():
        matches=[s for s in sources if s['task_id'].startswith(prefix+'_')]
        if len(matches)!=1:raise ValueError('Source task not unique: '+prefix)
        source=matches[0];folder=root/kind;path=folder/(kind+ext)
        put(kind+'/'+path.name,{'word_report':word_fixture,'ppt_cover':ppt_fixture,'web_sales':sheet_fixture}[kind]())
        if kind=='web_sales':
            for p in PRODUCTS:
                # Hidden values have to be read through the browser GUI at replay.
                page='<!doctype html><meta charset="utf-8"><title>Office Sales Lab</title><h1>'+html.escape(p['name'])+'</h1><p>Inspect the sales-data JSON using Developer Tools.</p><script id="sales-data" type="application/json">'+json.dumps(p)+'</script>'
                put(kind+'/site/'+p['id']+'.html',page.encode())
        differences={
            'word_report':['Windows Explorer → local Linux file manager; folder/blank DOCX prepared outside replay',
                'Word/WPS mixed source → WPS Writer; Acrobat → Firefox local PDF viewer',
                'English test prose replaces original text; heading/body editing, saving, PDF export and reading retained'],
            'ppt_cover':['Synthetic one-slide cover with five explicit objects replaces unavailable source PPT',
                'Repeated unsuccessful selections condensed; title resizing, deletions, footer/date and color edits retained',
                'Date box is duplicated and moved down; transparency/custom shade trials omitted; added save and file audit'],
            'web_sales':['Safari/online Ktown4u → Firefox/offline six-product pages; values are synthetic, not original sales',
                'Network inspector response search → Web Console inspection of embedded JSON, once per product',
                'Repeated search/correction attempts condensed; six browser/spreadsheet switches retained',
                'Ambiguous final formula and deletion of equals sign replaced by valid per-row multiplication and SUM; added save/audit']}
        spec=dict(version=1,kind=kind,title=title,source_task_id=source['task_id'],source_steps=source['source_steps'],
            document=str(path),folder=str(folder),pdf=str(folder/(kind+'.pdf')),exe=exe,
            steps=task_steps(kind),adaptations=differences[kind],profile=PROFILE,
            human_session_verified=False,gui_status='GUI_NOT_RUN',no_wss=True)
        js(kind+'/spec.json',spec)
        # Preserve text/code evidence for review only; never eval/exec it.
        js(kind+'/source.json',source)
        js(kind+'/scenario.json',dict(actions=[dict(type='wps_long_sequence',name=kind,app_key='WPS',spec=spec)]))
        entries.append(dict(kind=kind,title=title,source_task_id=source['task_id'],source_steps=source['source_steps'],
            adapted_steps=len(spec['steps']),scenario=str(folder/'scenario.json'),status='IMPLEMENTED_GUI_NOT_RUN',adaptations=differences[kind]))
    js('batch.json',dict(version=1,tasks=entries,gui_executed=False,wss_collected=False))
    put('README.md',README.encode())
    js('manifest.json',dict(complete=True,source_sha256=hashlib.sha256(SOURCE.read_bytes()).hexdigest(),artifacts=hashes))
    return entries


def audit_word(path):
    with zipfile.ZipFile(path) as z:
        doc=ET.fromstring(z.read('word/document.xml'));styles=ET.fromstring(z.read('word/styles.xml'))
    style_names={s.attrib.get('{'+NS['w']+'}styleId'):s.find('w:name',NS).attrib.get('{'+NS['w']+'}val','') for s in styles.findall('w:style',NS) if s.find('w:name',NS) is not None}
    observed=[]
    for p in doc.findall('.//w:body/w:p',NS):
        text=''.join(t.text or '' for t in p.findall('.//w:t',NS))
        if not text:continue
        style=p.find('w:pPr/w:pStyle',NS);sid=style.attrib['{'+NS['w']+'}val'] if style is not None else 'Normal'
        label=re.sub(r'\s','',style_names.get(sid,sid)).lower()
        observed.append((text,label))
    expected=[(t,s.lower()) for t,s in PARAGRAPHS]
    if observed!=expected:raise AssertionError('Word text or paragraph styles differ: '+repr(observed))
    return dict(paragraphs=len(observed),heading_styles=True)


def audit_ppt(path):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    p=Presentation(path)
    if len(p.slides)!=1:raise AssertionError('Slide count changed')
    shapes=list(p.slides[0].shapes);by_text={s.text:s for s in shapes if s.has_text_frame}
    required={'Why People Obey Authority','Mind and Society | Second Group Presentation','Reporting Group: Group 2','Report Date: 24.10.15'}
    if len(shapes)!=4 or set(by_text)!=required:raise AssertionError('PPT text/delete/duplicate result incorrect')
    for text,size in [('Why People Obey Authority',52),('Mind and Society | Second Group Presentation',20)]:
        runs=[r for p in by_text[text].text_frame.paragraphs for r in p.runs if r.text]
        if not runs or any(r.font.size is None or r.font.size.pt!=size for r in runs):raise AssertionError('Font size not preserved')
    date=by_text['Report Date: 24.10.15'];group=by_text['Reporting Group: Group 2']
    try:
        if date.fill.fore_color.rgb==RGBColor.from_string('FFD8B0') or date.line.color.rgb==RGBColor.from_string('000000'):raise AssertionError('Fill/outline unchanged')
    except (AttributeError,TypeError):raise AssertionError('Expected explicit solid fill and outline colors')
    if date.top<=group.top:raise AssertionError('Duplicated date box was not moved below group')
    return dict(objects=4,text=True,font_sizes=True,fill_and_outline_changed=True,date_moved=True)


def audit_sales(path):
    s=snapshot(path)
    if s['sheets']!=['Sales']:raise AssertionError('Unexpected worksheets')
    expected={'A1':'Product','B1':'Sales','C1':'Units','D1':'Product amount','A8':'Total'}
    values=dict(expected)
    for i,p in enumerate(PRODUCTS,2):
        for c,key in [('A','name'),('B','totalSaleAmount'),('C','units')]:expected[f'{c}{i}']=p[key];values[f'{c}{i}']=p[key]
        expected[f'D{i}']=f'=B{i}*C{i}';values[f'D{i}']=p['totalSaleAmount']*p['units']
    expected['D8']='=SUM(D2:D7)';values['D8']=sum(p['totalSaleAmount']*p['units'] for p in PRODUCTS)
    if s['formulas']!={'Sales':expected} or s['values']!={'Sales':values}:raise AssertionError('Sales input/formula/cached result mismatch')
    return dict(products=len(PRODUCTS),total=values['D8'],formulas_and_cached_values=True)


def audit_pdf(path):
    result=subprocess.run(['pdftotext',str(path),'-'],capture_output=True,text=True,check=True,timeout=20)
    text=' '.join(result.stdout.split())
    if not all(t in text for t,_ in PARAGRAPHS):raise AssertionError('PDF missing expected content')
    return dict(text_complete=True,bytes=path.stat().st_size)


class Desktop(EngineBackend):
    """Stateful GUI backend; OCR/profile misses fail rather than skip operations."""
    def __init__(self,engine,ctx,spec):
        super().__init__(engine,dict(path=spec['document'],title=Path(spec['document']).name,app_key='WPS',profile=spec['profile']),ctx)
        self.spec=spec;self.records={};self.browser_started=False;self.audit={}

    def raw_key(self,key):self.command(['xdotool','key','--clearmodifiers',key])
    def paste(self,text):
        self.e._set_clipboard(text.encode(),'UTF8_STRING',self.ctx);self.raw_key('ctrl+v');self.wait(.3)
    def click(self,x,y):self.command(['xdotool','mousemove',str(x),str(y),'click','1']);self.wait(.3)
    def launch(self,name,argv):
        guard().check(64*1024*1024)
        self.e.launch(dict(name=name,scope_name=name+'-'+Path(self.spec['folder']).parent.name,command=shlex.join(argv)),self.ctx)
    def window(self,title):
        app='WPS' if title==self.path.name else ('FILES' if title==Path(self.spec['folder']).name else 'FIREFOX')
        a=dict(title=title,name=title,app_key=app,timeout=30,strict_window_match=True)
        self.e.wait_window(a,self.ctx);wid=self.e.find_window(a,self.ctx)
        self.command(['wmctrl','-i','-r',wid,'-b','add,maximized_vert,maximized_horz'])
        self.command(['xdotool','windowactivate','--sync',wid]);self.wait(.5)
    def bind(self):
        self.window(self.path.name)
        deadline=time.monotonic()+35
        while time.monotonic()<deadline:
            try:super().ready();return
            except RuntimeError:self.wait(.5)
        raise RuntimeError('WPS embedded editor not available')
    def shape(self,index):
        self.bind();self.key('Escape');self.key('Escape')
        self.click(*self.spec['profile']['canvas_blank']);self.key('Escape')
        for _ in range(index):self.key('Tab')
    def browser(self,url,title='Office Sales Lab'):
        if not self.browser_started:
            folder=Path(self.spec['folder'])/'firefox-profile';guard().check();folder.mkdir()
            # Dedicated Firefox profile; no reuse of a personal browser session.
            self.launch('long_browser',['firefox','--no-remote','--profile',str(folder),url]);self.browser_started=True
        else:
            self.window('Mozilla Firefox');self.raw_key('ctrl+l');self.paste(url);self.raw_key('Return')
        self.window(title);self.wait(2)
    def save(self):
        guard().check(16*1024*1024);self.bind();self.key('ctrl+s');self.wait(2)
    def check(self,name,fn):
        deadline=time.monotonic()+15;last=None
        while time.monotonic()<deadline:
            try:self.audit[name]=fn();return
            except (AssertionError,OSError,zipfile.BadZipFile,subprocess.CalledProcessError) as exc:last=exc;self.wait(.5)
        raise AssertionError(f'{name}: {last}')

    def perform(self,step):
        op=step['op']
        if op=='open_document':
            self.launch('long_wps',['/opt/kingsoft/wps-office/office6/'+self.spec['exe'],str(self.path)])
            self.window(self.path.name);self.bind()
        elif op=='folder_browse':
            import shutil
            executable=next((shutil.which(n) for n in ('pcmanfm','thunar') if shutil.which(n)),None)
            if not executable:raise RuntimeError('No supported file manager installed')
            self.launch('long_files',[executable,self.spec['folder']]);self.window(Path(self.spec['folder']).name)
        elif op=='word_paragraph':
            self.bind();self.key('ctrl+End')
            # Calibrated against the 1920x1080 English WPS ribbon; audit the
            # saved paragraph style so coordinate drift cannot silently pass.
            self.click(358,50)
            self.click({'Heading1':878,'Heading2':948,'Normal':810}[step['style']],115)
            self.key('End')
            self.paste(step['text']);self.key('Return')
        elif op.startswith('shape_'):
            self.shape(step['index'])
            if op=='shape_delete':self.key('Delete')
            elif op=='shape_duplicate':
                self.key('ctrl+d')
                for _ in range(20):self.key('Down')
            elif op in ('shape_text','shape_font'):
                self.key('F2');self.key('ctrl+a')
                if op=='shape_text':self.paste(step['text'])
                else:
                    self.key('ctrl+shift+p');self.raw_key('ctrl+a');self.paste(str(step['size']));self.raw_key('Return')
                self.raw_key('Escape')
            elif op in ('shape_fill','shape_outline'):
                self.click_label(['Drawing Tools','绘图工具','Format','格式'],[0,25,1920,85])
                self.click_label(['Shape Fill','Fill','形状填充','填充'] if op=='shape_fill' else ['Shape Outline','Outline','形状轮廓','轮廓'],[0,65,1920,190])
                # Versioned palette coordinates require first-run calibration;
                # saved OOXML must show a real color change before success.
                self.click(*self.spec['profile']['fill_swatch' if op=='shape_fill' else 'line_swatch'])
                self.raw_key('Escape')
            else:raise ValueError(op)
        elif op=='save':self.save()
        elif op=='audit_word':self.check('word',lambda:audit_word(self.path))
        elif op=='audit_ppt':self.check('ppt',lambda:audit_ppt(self.path))
        elif op=='audit_sales':self.check('sales',lambda:audit_sales(self.path))
        elif op=='export_pdf':
            guard().check(16*1024*1024);self.bind()
            self.click_label(['Menu','文件','菜单'],[0,25,180,80])
            self.click_label(['Export to PDF','输出为PDF','导出为PDF','Export to PDF...'])
            point=self.locate(self.spec['profile']['pdf_path_label'])
            if point is None:raise RuntimeError('PDF destination field needs profile calibration')
            dx,dy=self.spec['profile']['pdf_path_offset'];self.click(point[0]+dx,point[1]+dy)
            self.raw_key('ctrl+a');self.paste(self.spec['folder'])
            self.click_label(['Export','Start Export','输出','开始输出','导出'])
            self.wait(2)
        elif op=='audit_pdf':self.check('pdf',lambda:audit_pdf(Path(self.spec['pdf'])))
        elif op=='read_pdf':
            self.browser(Path(self.spec['pdf']).as_uri(),Path(self.spec['pdf']).name)
            self.raw_key('ctrl+End');self.wait(2);self.raw_key('ctrl+Home');self.wait(2)
        elif op=='inspect_product':
            p=PRODUCTS[step['product_index']]
            self.browser((Path(self.spec['folder'])/'site'/(p['id']+'.html')).as_uri())
            self.raw_key('ctrl+shift+k');self.wait(2)
            # Type our fixed read-only DOM expression; never paste/evaluate source code.
            self.command(['xdotool','type','--clearmodifiers','--delay','10',"copy(document.querySelector('#sales-data').textContent)"])
            self.raw_key('Return');self.wait(1)
            raw=subprocess.run(['xclip','-selection','clipboard','-o'],capture_output=True,text=True,check=True,timeout=5).stdout
            observed=json.loads(raw)
            if observed!=p:raise AssertionError('Browser extraction mismatch or stale clipboard')
            self.records[step['product_index']]=observed
            self.raw_key('ctrl+shift+k')
        elif op=='record_product':
            p=self.records[step['product_index']];self.bind();self.select('A'+str(step['row']))
            self.paste(f'{p["name"]}\t{p["totalSaleAmount"]}\t{p["units"]}');self.key('Return')
        elif op=='sales_formulas':
            if len(self.records)!=len(PRODUCTS):raise AssertionError('Not all browser records collected')
            self.bind()
            for i in range(2,8):self.select(f'D{i}');self.paste(f'=B{i}*C{i}');self.key('Return')
            self.select('A8');self.paste('Total');self.key('Return')
            self.select('D8');self.paste('=SUM(D2:D7)');self.key('Return');self.key('F9')
        elif op=='close_apps':self.e.cleanup_tracked_processes(self.ctx)
        else:raise ValueError('Unknown long operation: '+op)
        self.wait(.5)


def validate(spec):
    if spec['kind'] not in TASKS:raise ValueError('Unknown scenario')
    if spec['steps']!=task_steps(spec['kind']):raise ValueError('Unexpected steps or modified executable parameters')
    g=guard();folder=g.path(str(Path(spec['folder']).relative_to(g.root)))
    for key in ('document','pdf'):
        path=g.path(str(Path(spec[key]).relative_to(g.root)))
        if path.parent!=folder:raise ValueError('Output outside task folder')
    if spec['exe']!=TASKS[spec['kind']][2]:raise ValueError('Unexpected executable')
    if not Path(spec['document']).is_file():raise FileNotFoundError(spec['document'])


def perform(action,ctx,engine):
    spec=action['spec'];validate(spec)
    if ctx.dry_run:
        for step in spec['steps']:engine.log('dry-run '+step['id']+' '+step['op'])
        return
    g=guard();relative=str(Path(spec['folder']).relative_to(g.root));result_path=relative+'/replay_result.json'
    if g.path(result_path).exists():raise FileExistsError('Fresh build required for every replay')
    # Refuse a workbook changed since preparation, including interrupted attempts.
    manifest=json.loads((Path(spec['folder']).parent/'manifest.json').read_text())
    key=spec['kind']+'/'+Path(spec['document']).name
    if hashlib.sha256(Path(spec['document']).read_bytes()).hexdigest()!=manifest['artifacts'][key]:raise ValueError('Fixture has changed; build fresh')
    backend=Desktop(engine,ctx,spec);events=[];error=''
    try:
        for step in spec['steps']:
            g.check(16*1024*1024)
            event=dict(id=step['id'],op=step['op'],start_ns=time.monotonic_ns(),status='RUNNING');events.append(event)
            engine.trace_marker_action(dict(type='trace_marker',operation_id=step['id'],event_type='OP_START',status='running',app_key='WPS'),ctx)
            backend.perform(step);event.update(end_ns=time.monotonic_ns(),status='PASS')
            engine.trace_marker_action(dict(type='trace_marker',operation_id=step['id'],event_type='OP_DONE',status='success',app_key='WPS'),ctx)
    except BaseException as exc:
        error=repr(exc)
        if events:events[-1].update(status='FAILED',end_ns=time.monotonic_ns(),error=error)
        if isinstance(exc,(KeyboardInterrupt,SystemExit)):raise
        raise engine.AutomationError(error) from exc
    finally:
        engine.cleanup_tracked_processes(ctx)
        g.write_json(result_path,dict(status='FAILED' if error else 'GUI_AUDIT_PASS',events=events,audits=backend.audit,
            error=error,source_task_id=spec['source_task_id'],adaptations=spec['adaptations'],wss_collected=False))


README='''# AgentNet 衍生长流程 v1

三条场景均为 IMPLEMENTED_GUI_NOT_RUN。这里交付实现、独立素材、逐阶段 trace 和任务级断言；未声称 GUI 实测通过，也未采集 WSS。

1. word_report：浏览目录 → 打开空白 DOCX → 输入 7 段标题/正文并设置样式 → 保存 → 检查 OOXML → WPS 导出 PDF → 检查 PDF 文本 → Firefox 阅读 → 关闭。
2. ppt_cover：打开封面 → 改标题与三次字号 → 删除标识和说明 → 改页脚 → 复制并下移汇报框 → 改日期 → 改填充与轮廓 → 调整页脚字号 → 保存并检查 → 关闭。
3. web_sales：打开表格 → 6 次 Firefox 商品页面/开发者控制台取值与 WPS 录入 → 6 行乘法和总和 → 保存并检查公式及缓存数值 → 关闭。

spec.json 记录每条适配差异和完整动作；source.json 保留源文本，仅供核对，绝不执行源 code。步数不是独立业务操作数。源轨迹不是已确认的长期用户会话。

使用 /home/lzx/Desktop/wss-ebpf-lightgbm/.venv/bin/python（含 openpyxl/python-pptx），PYTHONPATH 加入 test 和 wps_operation_dataset/.deps。先执行 scripts/build_wps_long_scenarios.py --output outputs/新目录，再用 app_automation.py 场景目录/scenario.json --dry-run 验证。实际运行用 scripts/run_wps_long_scenarios.py --batch 批次绝对路径 --run；默认只 dry-run。每次运行需要全新批次，避免覆盖已修改文件。runner 创建独立 Xvfb、openbox 和临时应用配置，结束时清理自己启动的进程；不启动 BPF/WSS。

首次 GUI 试跑需校准 profile：WPS 内嵌焦点、PPT Tab 对象顺序/空白画布位置、填充与边框调色板、PDF 导出菜单与保存路径字段。找不到标签或断言失败会报失败，不能跳过后冒充成功。固定 1920×1080、100% 缩放；当前标签支持英文/中文，不保证所有 WPS 版本通用。PDF 文件名沿用文档基本名，路径固定至当前任务目录。

输入素材由 Python 生成；运行结果只能由 GUI 写入。审计检查 Word 段落样式和文本、PDF 正文、PPT 文本/对象/字号/颜色变化/位置，以及表格输入、公式和缓存计算结果。保存和检查本身会影响后续 WSS 实验，trace 分开标记，当前不做内存测量。

网页是离线合成页面，无联网素材或账号；开发者工具读取 sales-data JSON，失败时不会退回直接读取本地文件作为浏览器结果。表格改用 Product/Sales/Units/Product amount 四列，预期总量 1336。未保留源末尾删除公式等号的行为，替换为可校验的计算。
'''
