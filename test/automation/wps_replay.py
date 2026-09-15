"""Independent spreadsheet operations and assertions for AgentNet replay.

Build/validate never starts a GUI. Runtime is invoked only through the existing
app_automation action dispatcher. Source pyautogui code is never executed.
"""
from collections import Counter
import csv
from datetime import datetime
import io
import json
from pathlib import Path
import re
import subprocess
import sys
import time
import zipfile
import xml.etree.ElementTree as ET

TEST_ROOT = Path(__file__).resolve().parents[1]
WORKSPACE = TEST_ROOT / 'wps_operation_dataset'
NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}


def guard():
    if str(WORKSPACE) not in sys.path:
        sys.path.insert(0, str(WORKSPACE))
    from wps_operation_dataset.storage import DiskGuard
    return DiskGuard(WORKSPACE)


def snapshot(path):
    """Read formulas and values separately; never write results with openpyxl."""
    import openpyxl
    f = openpyxl.load_workbook(path, data_only=False)
    v = openpyxl.load_workbook(path, data_only=True)
    try:
        return dict(sheets=f.sheetnames, formulas={s.title: {c.coordinate:c.value for row in s for c in row if c.value is not None} for s in f},
                    values={s.title: {c.coordinate:c.value for row in s for c in row if c.value is not None} for s in v})
    finally:
        f.close(); v.close()


def validate_spec(spec):
    if spec.get('operation') not in ('sort_ascending', 'recalculate_now'):
        raise ValueError('Unsupported operation')
    if not re.fullmatch(r'[A-Z]+[1-9][0-9]*', spec.get('cell','')):
        raise ValueError('Invalid cell')
    if spec['operation']=='sort_ascending':
        if spec.get('key_column') not in (3,4) or spec['cell'] != f"{chr(64+spec['key_column'])}2":
            raise ValueError('Sort key/cell mismatch')
        rows=spec.get('rows',[])
        if len(rows)<2 or any(len(r)!=4 for r in rows):raise ValueError('Expected non-empty four-column data')
        if spec.get('headers') != ['row_id','payload','key_c','key_d']:raise ValueError('Unexpected headers')
    elif spec['cell']!='B1':raise ValueError('NOW batch contract requires B1')


def audit_sort(path,spec):
    validate_spec(spec)
    snap=snapshot(path)
    if snap['sheets']!=['Data']:raise AssertionError('Sheet set changed')
    content=snap['formulas']['Data']; expected_headers=spec['headers']
    if [content.get(f'{chr(65+i)}1') for i in range(4)]!=expected_headers:raise AssertionError('Headers changed')
    rows=[[content.get(f'{chr(65+j)}{i}') for j in range(4)] for i in range(2,len(spec['rows'])+2)]
    expected=sorted(spec['rows'],key=lambda r:r[spec['key_column']-1])
    if rows!=expected:raise AssertionError('Rows not sorted or row associations changed')
    if len(content)!=4*(len(rows)+1):raise AssertionError('Unexpected content added/removed')
    return dict(ascending=True,whole_rows_preserved=True,headers_preserved=True,no_extra_cells=True)


def audit_now(path,baseline,started,ended):
    current=snapshot(path)
    if current['formulas']!=baseline['formulas'] or current['sheets']!=baseline['sheets']:
        raise AssertionError('Formula/content changed during recalculation')
    before=baseline['values']['Data'].get('B1'); after=current['values']['Data'].get('B1')
    if not isinstance(before,datetime) or not isinstance(after,datetime):raise AssertionError('NOW cache not a date/time')
    if after<=before:raise AssertionError('NOW cache did not advance')
    if not started-2 <= after.timestamp() <= ended+2:raise AssertionError('NOW value outside F9 observation interval')
    if current['formulas']['Data'].get('B1')!='=NOW()':raise AssertionError('NOW formula lost')
    other_before={k:v for k,v in baseline['values']['Data'].items() if k!='B1'}
    other_after={k:v for k,v in current['values']['Data'].items() if k!='B1'}
    if other_before!=other_after:raise AssertionError('Other cached cells changed')
    return dict(formula_preserved=True,cache_advanced=True,cache_in_time_interval=True,other_cells_preserved=True)


def execute_semantic(spec,backend,path,timeout=25):
    """Backend is injectable for contract tests; no assertions are faked by it."""
    validate_spec(spec)
    backend.ready()
    # Persist baseline outside the source operation window. This also guards
    # against automatic recalculation on initial document load.
    backend.save()
    baseline=snapshot(path)
    if spec['operation']=='sort_ascending':
        expected={f'{chr(65+j)}{i+2}':v for i,r in enumerate(spec['rows']) for j,v in enumerate(r)}
        expected.update({f'{chr(65+j)}1':v for j,v in enumerate(spec['headers'])})
        if baseline['formulas']!={'Data':expected}:raise AssertionError('Sort fixture is not pristine')
    else:
        if baseline['formulas']['Data'].get('B1')!='=NOW()':raise AssertionError('Initial NOW formula missing')
        displayed_before=backend.read_cell_time('B1')
        backend.wait(1.5)  # Ensure second-resolution date/time can advance.
    started=time.time()
    backend.event('OP_START',spec['task_id'])
    try:
        backend.select(spec['cell'])
        if spec['operation']=='sort_ascending':
            backend.sort_ascending()
        else:
            backend.key('Menu');backend.wait(.3);backend.key('Escape')
            backend.select('B1');backend.key('F9');backend.wait(1)
            backend.select('A1');backend.select('B1')
    except Exception:
        backend.event('OP_FAILED',spec['task_id']);raise
    finally:
        ended=time.time()
    backend.event('OP_ACTIONS_RETURNED',spec['task_id'])
    if spec['operation']=='recalculate_now':
        # Read the live result BEFORE audit-only saving. A save-triggered
        # recalculation alone must not make a non-working F9 operation pass.
        displayed_after=backend.read_cell_time('B1')
        if displayed_after<=displayed_before or not started-2<=displayed_after.timestamp()<=time.time()+2:
            backend.event('OP_FAILED',spec['task_id'])
            raise AssertionError('Live NOW value did not advance before audit save')
    # Audit-only save outside source window. Never call openpyxl.save here.
    backend.save()
    deadline=time.monotonic()+timeout
    while True:
        try:
            audit=audit_sort(path,spec) if spec['operation']=='sort_ascending' else audit_now(path,baseline,started,ended)
            if spec['operation']=='recalculate_now':audit['live_value_advanced_before_save']=True
            backend.event('OP_AUDIT_PASS',spec['task_id'])
            return dict(status='GUI_AUDIT_PASS',audit=audit,started=started,ended=ended,
                        no_wss_collected=True,source_navigation_adapted=True)
        except (AssertionError,OSError,zipfile.BadZipFile) as exc:
            if time.monotonic()>=deadline:
                backend.event('OP_FAILED',spec['task_id']);raise RuntimeError('Postcondition failed: '+str(exc)) from exc
            backend.wait(.5)


class EngineBackend:
    """Uses the existing dispatcher/clipboard primitives plus embedded focus/OCR."""
    def __init__(self,engine,action,ctx):
        self.e,self.a,self.ctx=engine,action,ctx
        self.path=Path(action['path']);self.wid=None

    def command(self,args):
        return self.e.run(args,self.ctx)

    def ready(self):
        a={**self.a,'app_key':'WPS','title':self.path.name}
        root=self.e.best_window_candidate(a,'WPS')
        if root is None or root.mapped_app!='WPS' or self.path.name not in root.title:
            raise RuntimeError('Target WPS workbook window not found')
        self.command(['xdotool','windowactivate','--sync',root.window_id])
        result=self.command(['xdotool','search','--onlyvisible','--name',re.escape(self.path.name)])
        for wid in result.stdout.splitlines():
            if int(wid)==int(root.window_id,16):continue
            info=self.e.read_window_info(wid)
            if info.mapped_app!='WPS' or info.width<800 or info.height<400:continue
            if root.cgroup_path and info.cgroup_path!=root.cgroup_path:continue
            self.wid=wid;self.command(['xdotool','windowfocus',wid]);return
        raise RuntimeError('Owned embedded spreadsheet editor not found')

    def key(self,key):
        if self.wid is None:raise RuntimeError('Editor not bound')
        self.command(['xdotool','windowfocus',self.wid,'key','--clearmodifiers',key])

    def select(self,ref):
        match=re.fullmatch(r'([A-Z]+)([1-9][0-9]*)',ref)
        if not match:raise ValueError('Invalid cell')
        column=0
        for char in match[1]:column=column*26+ord(char)-64
        row=int(match[2]);self.key('ctrl+Home')
        for key,count in [('Right',column-1),('Down',row-1)]:
            if count:self.command(['xdotool','windowfocus',self.wid,'key','--clearmodifiers','--repeat',str(count),'--delay','20',key])

    def wait(self,seconds):time.sleep(seconds)

    def read_cell_time(self,reference):
        self.select(reference);self.key('ctrl+c');self.wait(.3)
        raw=subprocess.run(['xclip','-selection','clipboard','-o'],capture_output=True,text=True,check=True,timeout=5).stdout.strip()
        self.key('Escape')
        try:return datetime.fromisoformat(raw.replace('/','-'))
        except ValueError as exc:raise RuntimeError('Live NOW cell text is not an unambiguous ISO date/time: '+repr(raw)) from exc

    def ocr(self):
        capture=subprocess.run(['import','-silent','-window','root','png:-'],capture_output=True,check=True,timeout=10)
        result=subprocess.run(['tesseract','stdin','stdout','-l','eng+chi_sim','--psm','11','tsv'],input=capture.stdout,capture_output=True,check=True,timeout=30)
        return list(csv.DictReader(io.StringIO(result.stdout.decode()),delimiter='\t',quoting=csv.QUOTE_NONE))

    def locate(self,aliases,region=None):
        rows=[r for r in self.ocr() if r.get('text','').strip()]
        norm=lambda t:re.sub(r'[^a-z0-9\u4e00-\u9fff]','',t.lower())
        for alias in aliases:
            found=set()
            for i,r in enumerate(rows):
                for n in range(1,5):
                    group=rows[i:i+n]
                    if len({(t['block_num'],t['par_num'],t['line_num']) for t in group})!=1:break
                    if norm(''.join(t['text'] for t in group))!=norm(alias):continue
                    x=min(int(t['left']) for t in group);y=min(int(t['top']) for t in group)
                    right=max(int(t['left'])+int(t['width']) for t in group);bottom=max(int(t['top'])+int(t['height']) for t in group)
                    pt=((x+right)//2,(y+bottom)//2)
                    if region is None or region[0]<=pt[0]<=region[2] and region[1]<=pt[1]<=region[3]:found.add(pt)
            if len(found)==1:return next(iter(found))
            if len(found)>1:raise RuntimeError('Ambiguous UI text: '+alias)
        return None

    def click_label(self,aliases,region=None):
        point=self.locate(aliases,region)
        if point is None:raise RuntimeError('UI label unavailable; profile needs calibration: '+str(aliases))
        self.command(['xdotool','mousemove',str(point[0]),str(point[1]),'click','1']);self.wait(.5)

    def sort_ascending(self):
        profile=self.a['profile']
        self.click_label(profile['data_tab'],[0,0,1920,190])
        if self.locate(profile['ascending'],[0,0,1920,240]) is None:
            # This WPS build places ascending inside the labelled Sort menu.
            self.click_label(['Sort','排序'],[0,70,600,150])
        self.click_label(profile['ascending'],[0,0,1920,500])
        point=self.locate(profile['expand_selection'])
        if point:
            self.command(['xdotool','mousemove',str(point[0]),str(point[1]),'click','1'])
            self.click_label(profile['sort_confirm'])
        self.wait(1)
        # If WPS did not expand rows or picked the wrong column, workbook audit
        # rejects it. No extra filter/copy/formula actions are added.

    def save(self):
        guard().check()
        before=self.path.stat().st_mtime_ns
        self.key('ctrl+s')
        deadline=time.monotonic()+15
        while time.monotonic()<deadline:
            self.wait(.5)
            try:
                snapshot(self.path)
                if self.path.stat().st_mtime_ns>before:return
            except (OSError,zipfile.BadZipFile):pass
        # Initial unmodified save may leave mtime unchanged; content checks still
        # reject stale output after a real sort/recalculation.
        snapshot(self.path)

    def event(self,event,task_id):
        self.e.trace_marker_action(dict(type='trace_marker',event_type=event,operation_id=task_id,
            status=event.lower(),app_key='WPS'),self.ctx)


def perform(action,ctx,engine):
    spec=action['spec'];validate_spec(spec)
    path=Path(action['path']).resolve()
    g=guard();g.path(str(path.relative_to(WORKSPACE)))
    if not path.is_file():raise FileNotFoundError(path)
    if ctx.dry_run:
        engine.log('dry-run: semantic WPS task '+spec['task_id']+'; GUI_NOT_RUN')
        return
    g.check()
    relative=str(path.parent.relative_to(WORKSPACE)/'replay_result.json')
    if g.path(relative).exists():raise FileExistsError('Result exists; prepare a fresh task workbook')
    try:
        result=execute_semantic(spec,EngineBackend(engine,action,ctx),path)
        result['task_id']=spec['task_id'];g.write_json(relative,result)
    except Exception as exc:
        if not g.path(relative).exists():g.write_json(relative,dict(status='FAILED',task_id=spec['task_id'],error=str(exc),no_wss_collected=True))
        raise


def workbook_bytes(spec):
    import openpyxl
    from openpyxl.workbook.properties import CalcProperties
    validate_spec(spec)
    book=openpyxl.Workbook();sheet=book.active;sheet.title='Data'
    if spec['operation']=='sort_ascending':
        sheet.append(spec['headers'])
        for row in spec['rows']:sheet.append(row)
    else:
        sheet['A1']='当前时间';sheet['B1']='=NOW()';sheet['B1'].number_format='yyyy-mm-dd hh:mm:ss'
        book.calculation=CalcProperties(calcMode='manual',fullCalcOnLoad=False,forceFullCalc=False,calcOnSave=False)
    buf=io.BytesIO();book.save(buf);book.close()
    if spec['operation']=='sort_ascending':return buf.getvalue()
    # Valid stale cached value: only WPS may update it during future replay.
    out=io.BytesIO()
    with zipfile.ZipFile(buf) as src,zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data=src.read(item.filename)
            if item.filename=='xl/worksheets/sheet1.xml':
                xml=ET.fromstring(data);cell=xml.find('.//s:c[@r="B1"]',NS)
                value=cell.find('s:v',NS);value.text='36526'
                data=ET.tostring(xml,encoding='utf-8',xml_declaration=True)
            dst.writestr(item,data)
    return out.getvalue()
