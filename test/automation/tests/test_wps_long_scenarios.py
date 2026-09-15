import io
import json
from pathlib import Path
import socket
import subprocess
import xml.etree.ElementTree as ET
import zipfile

import pytest

from automation import app_automation as engine
from automation import wps_long_scenarios as long


@pytest.fixture
def local_guard(tmp_path,monkeypatch):
    from wps_operation_dataset.storage import DiskGuard
    g=DiskGuard(tmp_path)
    monkeypatch.setattr(long,'guard',lambda:g)
    return g


def test_build_and_entire_batch_dry_run_without_processes_or_network(local_guard,monkeypatch):
    def deny(*a,**kw):pytest.fail('Unexpected process/network call during build or dry-run')
    monkeypatch.setattr(subprocess,'Popen',deny);monkeypatch.setattr(socket.socket,'connect',deny)
    entries=long.build('batch')
    assert [e['source_steps'] for e in entries]==[47,59,118]
    for entry in entries:
        s=engine.load_scenario(Path(entry['scenario']))
        for action in s.actions:engine.ACTION_HANDLERS[action['type']](action,engine.Context(dry_run=True))
        assert not Path(entry['scenario']).with_name('replay_result.json').exists()
    with pytest.raises(FileExistsError):long.build('batch')


def test_source_and_adapted_chains_are_distinct(local_guard):
    entries=long.build('batch')
    for entry in entries:
        spec=json.loads(Path(entry['scenario']).with_name('spec.json').read_text())
        assert spec['adaptations'] and spec['gui_status']=='GUI_NOT_RUN'
        assert len({s['id'] for s in spec['steps']})==len(spec['steps'])
        assert spec['steps'][-1]['op']=='close_apps'
        assert not spec['human_session_verified']
    web=long.task_steps('web_sales')
    assert [s['op'] for s in web[1:13]]==['inspect_product','record_product']*6


def test_word_audit_rejects_blank_and_wrong_styles(tmp_path):
    p=tmp_path/'x.docx';p.write_bytes(long.word_fixture())
    with pytest.raises(AssertionError):long.audit_word(p)
    original=p.read_bytes()
    def set_paragraphs(wrong=False):
        with zipfile.ZipFile(io.BytesIO(original)) as z:parts={n:z.read(n) for n in z.namelist()}
        ns=long.NS['w'];tag=lambda n:'{'+ns+'}'+n
        doc=ET.fromstring(parts['word/document.xml']);body=doc.find('w:body',long.NS)
        for text,style in long.PARAGRAPHS:
            para=ET.Element(tag('p'));prop=ET.SubElement(para,tag('pPr'))
            ET.SubElement(prop,tag('pStyle'),{tag('val'):'Normal' if wrong else style})
            ET.SubElement(ET.SubElement(para,tag('r')),tag('t')).text=text
            body.insert(len(body)-1,para)
        parts['word/document.xml']=ET.tostring(doc)
        with zipfile.ZipFile(p,'w') as z:
            for n,v in parts.items():z.writestr(n,v)
    set_paragraphs();assert long.audit_word(p)['heading_styles']
    set_paragraphs(True)
    with pytest.raises(AssertionError):long.audit_word(p)


def test_ppt_audit_checks_content_font_color_and_position(tmp_path):
    from pptx import Presentation
    from pptx.util import Pt, Inches
    from pptx.dml.color import RGBColor
    from copy import deepcopy
    p=tmp_path/'x.pptx';p.write_bytes(long.ppt_fixture())
    with pytest.raises(AssertionError):long.audit_ppt(p)
    ppt=Presentation(p);sh=ppt.slides[0].shapes
    for index in [2,1]:
        el=sh[index]._element;el.getparent().remove(el)
    clone=deepcopy(sh[2]._element);clone.xpath('.//p:cNvPr')[0].set('id','99')
    sh._spTree.insert_element_before(clone,'p:extLst')
    for i,text in enumerate(['Why People Obey Authority','Mind and Society | Second Group Presentation','Reporting Group: Group 2','Report Date: 24.10.15']):sh[i].text=text
    sh[0].text_frame.paragraphs[0].runs[0].font.size=Pt(52)
    sh[1].text_frame.paragraphs[0].runs[0].font.size=Pt(20)
    sh[3].top=sh[2].top+Inches(1)
    sh[3].fill.solid();sh[3].fill.fore_color.rgb=RGBColor.from_string('CC3333')
    sh[3].line.color.rgb=RGBColor.from_string('CC3333');ppt.save(p)
    assert long.audit_ppt(p)['date_moved']
    sh[3].line.color.rgb=RGBColor.from_string('000000');ppt.save(p)
    with pytest.raises(AssertionError,match='unchanged'):long.audit_ppt(p)


def test_sales_rejects_formula_strings_without_calculation(tmp_path):
    import openpyxl
    p=tmp_path/'x.xlsx';p.write_bytes(long.sheet_fixture());b=openpyxl.load_workbook(p);s=b.active
    for i,product in enumerate(long.PRODUCTS,2):
        s.append([product['name'],product['totalSaleAmount'],product['units'],f'=B{i}*C{i}'])
    s['A8']='Total';s['D8']='=SUM(D2:D7)';b.save(p);b.close()
    with pytest.raises(AssertionError):long.audit_sales(p)
    # Test fixture only: populate expected caches, then verify audit rejects a
    # wrong result even if the displayed formula text is unchanged.
    ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    with zipfile.ZipFile(p) as z:parts={n:z.read(n) for n in z.namelist()}
    xml=ET.fromstring(parts['xl/worksheets/sheet1.xml'])
    expected={f'D{i}':p['totalSaleAmount']*p['units'] for i,p in enumerate(long.PRODUCTS,2)}
    expected['D8']=sum(expected.values())
    def cached(wrong=False):
        for cell in xml.findall('.//s:c',ns):
            if cell.attrib['r'] in expected:cell.find('s:v',ns).text=str(expected[cell.attrib['r']]+int(wrong))
        parts['xl/worksheets/sheet1.xml']=ET.tostring(xml)
        with zipfile.ZipFile(p,'w') as z:
            for n,v in parts.items():z.writestr(n,v)
    cached();assert long.audit_sales(p)['total']==1336
    cached(True)
    with pytest.raises(AssertionError):long.audit_sales(p)


def test_browser_extraction_rejects_stale_clipboard(monkeypatch):
    spec=dict(document='/tmp/test.xlsx',folder='/tmp',profile=long.PROFILE)
    b=long.Desktop(engine,engine.Context(dry_run=True),spec)
    b.browser=lambda *a:None;b.raw_key=lambda *a:None;b.wait=lambda *a:None;b.command=lambda *a:None
    monkeypatch.setattr(subprocess,'run',lambda *a,**kw:type('Result',(),{'stdout':json.dumps(long.PRODUCTS[1])})())
    with pytest.raises(AssertionError,match='stale clipboard'):b.perform(dict(op='inspect_product',product_index=0))
    assert not b.records


def test_invalid_parameters_and_paths_fail_before_gui(local_guard):
    entries=long.build('batch');spec=json.loads(Path(entries[0]['scenario']).with_name('spec.json').read_text())
    long.validate(spec)
    spec['steps'][0]['op']='shell'
    with pytest.raises(ValueError):long.validate(spec)
    spec['steps']=long.task_steps('word_report');spec['pdf']='/tmp/escape.pdf'
    with pytest.raises(ValueError):long.validate(spec)


def test_failure_preserves_step_error_and_cleans_processes(local_guard,monkeypatch):
    entry=long.build('batch')[0];spec=json.loads(Path(entry['scenario']).with_name('spec.json').read_text())
    class Backend:
        audit={}
        def __init__(self,*args):pass
        def perform(self,step):raise RuntimeError('OCR label missing')
    monkeypatch.setattr(long,'Desktop',Backend)
    called=[];monkeypatch.setattr(engine,'cleanup_tracked_processes',lambda ctx:called.append('cleanup'))
    monkeypatch.setattr(engine,'trace_marker_action',lambda *a:None)
    with pytest.raises(engine.AutomationError,match='OCR label missing'):
        long.perform(dict(spec=spec),engine.Context(dry_run=False),engine)
    result=json.loads(Path(spec['folder'],'replay_result.json').read_text())
    assert result['status']=='FAILED' and result['events'][0]['status']=='FAILED' and called
    with pytest.raises(FileExistsError):long.perform(dict(spec=spec),engine.Context(dry_run=False),engine)
