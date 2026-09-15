from datetime import datetime,timedelta
import io
import json
from pathlib import Path
import socket
import subprocess
import time
from types import SimpleNamespace
import zipfile
import xml.etree.ElementTree as ET

import openpyxl
import pytest

from automation import wps_replay as replay
from automation import app_automation as engine
from automation.scripts.build_agentnet_replay_batch import TASKS


def sort_spec(column=3):
    return dict(task_id='test',operation='sort_ascending',cell=chr(64+column)+'2',key_column=column,
        headers=['row_id','payload','key_c','key_d'],rows=[['a','first',30,100],['b','second',10,300],['c','third',20,200]])


def write_fixture(tmp_path,spec):
    path=tmp_path/'book.xlsx';path.write_bytes(replay.workbook_bytes(spec));return path


def edit_rows(path,rows):
    b=openpyxl.load_workbook(path)
    for i,row in enumerate(rows,2):
        for j,v in enumerate(row,1):b.active.cell(i,j,v)
    b.save(path);b.close()


def set_cached_time(path,dt):
    from openpyxl.utils.datetime import to_excel
    data=io.BytesIO()
    with zipfile.ZipFile(path) as src,zipfile.ZipFile(data,'w') as dst:
        for item in src.infolist():
            content=src.read(item.filename)
            if item.filename=='xl/worksheets/sheet1.xml':
                root=ET.fromstring(content);root.find('.//s:c[@r="B1"]/s:v',replay.NS).text=str(to_excel(dt));content=ET.tostring(root)
            dst.writestr(item,content)
    path.write_bytes(data.getvalue())


@pytest.mark.parametrize('column',[3,4])
def test_sort_audit_checks_whole_rows_and_correct_column(tmp_path,column):
    spec=sort_spec(column);p=write_fixture(tmp_path,spec)
    with pytest.raises(AssertionError):replay.audit_sort(p,spec)
    expected=sorted(spec['rows'],key=lambda r:r[column-1]);edit_rows(p,expected)
    assert replay.audit_sort(p,spec)['whole_rows_preserved']
    broken=[list(r) for r in spec['rows']]
    for i,v in enumerate(sorted(r[column-1] for r in broken)):broken[i][column-1]=v
    edit_rows(p,broken)
    with pytest.raises(AssertionError):replay.audit_sort(p,spec)


def test_sort_rejects_extra_business_effects(tmp_path):
    s=sort_spec();p=write_fixture(tmp_path,s);edit_rows(p,sorted(s['rows'],key=lambda r:r[2]))
    b=openpyxl.load_workbook(p);b.active['M1']='extra formula';b.save(p);b.close()
    with pytest.raises(AssertionError,match='Unexpected'):replay.audit_sort(p,s)


def test_now_fixture_and_stale_cache_rejection(tmp_path):
    spec=dict(task_id='now',operation='recalculate_now',cell='B1');p=write_fixture(tmp_path,spec)
    baseline=replay.snapshot(p)
    assert baseline['formulas']['Data']['B1']=='=NOW()'
    assert baseline['values']['Data']['B1'].year==2000
    b=openpyxl.load_workbook(p);assert b.calculation.calcMode=='manual' and not b.calculation.calcOnSave;b.close()
    with pytest.raises(AssertionError,match='did not advance'):replay.audit_now(p,baseline,time.time(),time.time())
    stamp=datetime.now();set_cached_time(p,stamp)
    assert replay.audit_now(p,baseline,stamp.timestamp()-1,stamp.timestamp()+1)['cache_advanced']
    with pytest.raises(AssertionError,match='interval'):replay.audit_now(p,baseline,stamp.timestamp()+60,stamp.timestamp()+90)


def test_operation_window_excludes_audit_saves(tmp_path):
    spec=sort_spec();p=write_fixture(tmp_path,spec);calls=[]
    class Fake:
        saves=0
        def ready(self):calls.append('ready')
        def save(self):
            calls.append('save');self.saves+=1
            if self.saves==2:edit_rows(p,sorted(spec['rows'],key=lambda r:r[2]))
        def select(self,ref):calls.append('select:'+ref)
        def sort_ascending(self):calls.append('sort')
        def event(self,event,task):calls.append(event)
        def wait(self,seconds):pass
    result=replay.execute_semantic(spec,Fake(),p,timeout=0)
    assert result['status']=='GUI_AUDIT_PASS'
    assert calls==['ready','save','OP_START','select:C2','sort','OP_ACTIONS_RETURNED','save','OP_AUDIT_PASS']


def test_failed_backend_never_reports_audit_pass(tmp_path):
    spec=sort_spec();p=write_fixture(tmp_path,spec);events=[]
    class Fake:
        def ready(self):pass
        def save(self):pass
        def select(self,r):pass
        def sort_ascending(self):raise RuntimeError('missing OCR label')
        def event(self,e,t):events.append(e)
    with pytest.raises(RuntimeError,match='OCR'):replay.execute_semantic(spec,Fake(),p,timeout=0)
    assert 'OP_FAILED' in events and 'OP_AUDIT_PASS' not in events


def test_now_cannot_pass_due_only_to_saving(tmp_path):
    spec=dict(task_id='now',operation='recalculate_now',cell='B1');p=write_fixture(tmp_path,spec)
    class Fake:
        saves=0
        def ready(self):pass
        def save(self):self.saves+=1
        def select(self,ref):pass
        def key(self,key):pass
        def wait(self,seconds):pass
        def event(self,event,task):pass
        def read_cell_time(self,ref):return datetime(2000,1,1)
    backend=Fake()
    with pytest.raises(AssertionError,match='before audit save'):
        replay.execute_semantic(spec,backend,p,timeout=0)
    assert backend.saves==1


def test_now_positive_live_and_persisted_checks(tmp_path):
    spec=dict(task_id='now',operation='recalculate_now',cell='B1');p=write_fixture(tmp_path,spec)
    class Fake:
        saves=0
        reads=0
        def ready(self):pass
        def save(self):
            self.saves+=1
            if self.saves==2:set_cached_time(p,datetime.now())
        def select(self,ref):pass
        def key(self,key):pass
        def wait(self,seconds):pass
        def event(self,event,task):pass
        def read_cell_time(self,ref):
            self.reads+=1
            return datetime(2000,1,1) if self.reads==1 else datetime.now()
    assert replay.execute_semantic(spec,Fake(),p,timeout=0)['audit']['live_value_advanced_before_save']


def test_ocr_rejects_ambiguous_labels():
    b=replay.EngineBackend(None,{'path':'x'},None)
    b.ocr=lambda:[dict(text='Data',left=str(x),top='10',width='40',height='15',block_num=str(x),par_num='1',line_num='1') for x in (10,200)]
    with pytest.raises(RuntimeError,match='Ambiguous'):b.locate(['Data'])


@pytest.mark.parametrize('menu_required',[False,True])
def test_sort_reveals_menu_only_when_ascending_is_hidden(menu_required):
    from automation.scripts.build_agentnet_replay_batch import PROFILE
    b=replay.EngineBackend(None,{'path':'x','profile':PROFILE},None)
    calls=[]
    def locate(aliases,region=None):
        if aliases==PROFILE['ascending'] and not menu_required:return (100,100)
        return None
    b.locate=locate
    b.click_label=lambda aliases,region=None:calls.append(aliases)
    b.wait=lambda seconds:None
    b.sort_ascending()
    expected=[PROFILE['data_tab']]
    if menu_required:expected.append(['Sort','排序'])
    assert calls==expected+[PROFILE['ascending']]


def test_batch_build_and_whole_scenario_dry_run_offline(tmp_path,monkeypatch):
    from automation.scripts import build_agentnet_replay_batch as builder
    from wps_operation_dataset.storage import DiskGuard
    def deny(*a,**k):pytest.fail('Process/network attempted during preparation/dry-run')
    monkeypatch.setattr(subprocess,'Popen',deny);monkeypatch.setattr(socket.socket,'connect',deny)
    monkeypatch.setattr(builder,'guard',lambda:DiskGuard(tmp_path))
    monkeypatch.setattr(replay,'guard',lambda:DiskGuard(tmp_path));monkeypatch.setattr(replay,'WORKSPACE',tmp_path)
    entries=builder.build('batch')
    assert len(entries)==3 and all(t['source_step_coverage_complete'] for t in entries)
    for t in entries:
        assert t['runtime_status']=='GUI_NOT_RUN'
        scenario=engine.load_scenario(Path(t['scenario']))
        ctx=engine.Context(dry_run=True)
        for action in scenario.actions:engine.ACTION_HANDLERS[action['type']](action,ctx)
        assert not Path(t['fixture']).with_name('replay_result.json').exists()


def test_action_registered_and_dry_run_does_not_start_gui(tmp_path,monkeypatch):
    def deny(*a,**k):pytest.fail('GUI/network action attempted')
    monkeypatch.setattr(subprocess,'Popen',deny);monkeypatch.setattr(socket.socket,'connect',deny)
    from wps_operation_dataset.storage import DiskGuard
    monkeypatch.setattr(replay,'WORKSPACE',tmp_path);monkeypatch.setattr(replay,'guard',lambda:DiskGuard(tmp_path))
    s=sort_spec();p=write_fixture(tmp_path,s)
    assert 'wps_replay_operation' in engine.ACTION_HANDLERS
    engine.wps_replay_action(dict(path=str(p),spec=s),engine.Context(dry_run=True))
    assert not (tmp_path/'replay_result.json').exists()


def test_original_tasks_full_source_step_mapping():
    assert len(TASKS)==3
    for task in TASKS:
        assert [i for group in task['source_mapping'] for i in group['positions']]==list(range(7))
